#!/usr/bin/env python3
"""
build_outreach_system.py - GardenSuite Automated Outreach Filtering & Generation System

Ingests Master Registries for West Bengal (396) and Assam (736).
Applies:
1. Corporate Purge (Conglomerates, Public Listed, Multi-garden groups >2, Current Clients)
2. Scale & Operating Profile Filters
3. Tech-Leakage Scorer (detects paper/Tally/consumer tech)
4. Decision-Maker & Channel Allocation (WhatsApp, Postal, Email)
5. Dynamic Engagement Hook Generator (Zero product pitch, diagnostic questions)

Outputs:
- marketing/outreach/filtered/wb_filtered_candidates.csv
- marketing/outreach/filtered/assam_filtered_candidates.csv
- marketing/outreach/send_ready/wb_final_send_ready.csv
- marketing/outreach/send_ready/wb_blf_grower_outreach.csv
- marketing/outreach/send_ready/wb_siliguri_local_outreach.csv
- marketing/outreach/send_ready/assam_final_send_ready.csv
"""

import os
import csv
import re
from collections import Counter

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))
WB_MASTER = os.path.join(REPO_ROOT, "marketing/outreach/north-bengal-intelligence/data/North_Bengal_Tea_Master_Registry.csv")
ASSAM_MASTER = os.path.join(REPO_ROOT, "marketing/outreach/assam-intelligence/data/Assam_Tea_Master_Registry.csv")

FILTERED_DIR = os.path.join(REPO_ROOT, "marketing/outreach/filtered")
SEND_READY_DIR = os.path.join(REPO_ROOT, "marketing/outreach/send_ready")

# Existing GardenSuite clients from AGENTS.md - MUST BE STRICTLY EXCLUDED
CURRENT_CLIENTS = {
    "SIMULBARIE", "LONGVIEW", "RHEABARI", "MOGULKATA", "RAHIMPUR", "DEBPARA",
    "KURTI", "LOOKSAN", "SUBHASINI", "CHOIBARI", "CHAPAR", "DOOLAHAT", "ATAL",
    "THANJHORA", "NAXALBARI", "PAHARGOOMIAH", "KAMALPUR", "TINBIGHA", "CHANDAN",
    "HIMALAYAN AGRO", "HARISHPUR", "BAGRODIA", "BAGRODIA PLANTATIONS"
}

# Known corporate conglomerates across North Bengal and Assam
CONGLOMERATE_KEYWORDS = [
    "GOODRICKE", "DUNCANS", "MCLEOD RUSSEL", "AMALGAMATED PLANTATIONS",
    "TATA CONSUMER", "ANDREW YULE", "APEEJAY", "WARREN TEA", "ROSSELL",
    "DHUNSERI", "ASSAM COMPANY", "JAYSHREE TEA", "DIANA TEA", "GILLANDERS",
    "RYDAK SYNDICATE", "LUXMI TEA", "WILLIAMSON MAGOR", "CAMELLIA", "HASIMARA",
    "BUXA DOOARS", "GOVERNMENT OF INDIA", "BRS GROUP", "CHAMONG TEE", "CHAMONG"
]

CONSUMER_EMAIL_DOMAINS = {
    "gmail.com", "yahoo.com", "yahoo.co.in", "yahoo.in", "rediffmail.com",
    "vsnl.net", "vsnl.com", "bsnl.in", "sancharnet.in", "sifi.com",
    "indiatimes.com", "mail.com", "outlook.com", "hotmail.com"
}

def clean_text(val):
    if val is None:
        return ""
    val = str(val).strip()
    return "" if val == "-" else val

def parse_hectares(scale_str):
    if not scale_str:
        return None
    match = re.search(r"([\d,.]+)\s*(?:Ha|Hectares?|ha)", scale_str, re.IGNORECASE)
    if match:
        try:
            return float(match.group(1).replace(",", ""))
        except ValueError:
            return None
    return None

def is_mobile_number(phone_str):
    if not phone_str:
        return False
    digits = re.sub(r"[^\d]", "", phone_str)
    # Check for 10 digit Indian mobile starting with 6, 7, 8, 9
    if len(digits) == 10 and digits[0] in "6789":
        return True
    # Check for 12 digits starting with 91 followed by 6, 7, 8, 9
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return True
    # Check for 11 digits starting with 0 followed by 6, 7, 8, 9
    if len(digits) == 11 and digits.startswith("0") and digits[1] in "6789":
        return True
    return False

def clean_phone_number(phone_str):
    if not phone_str:
        return ""
    digits = re.sub(r"[^\d]", "", phone_str)
    if len(digits) == 10 and digits[0] in "6789":
        return digits
    if len(digits) == 12 and digits.startswith("91") and digits[2] in "6789":
        return digits[2:]
    if len(digits) == 11 and digits.startswith("0") and digits[1] in "6789":
        return digits[1:]
    return phone_str

def is_corporate(company_name, garden_name, company_counts):
    comp_upper = company_name.upper()
    gard_upper = garden_name.upper()

    # Check against current clients
    for client in CURRENT_CLIENTS:
        if client in comp_upper or client in gard_upper:
            return True, f"Current GardenSuite client ({client})"

    # Check against known conglomerates
    for cong in CONGLOMERATE_KEYWORDS:
        if cong in comp_upper or cong in gard_upper:
            return True, f"Conglomerate keyword match ({cong})"

    # Check public company marker: "LIMITED" without "PRIVATE"
    if "LIMITED" in comp_upper or "LTD" in comp_upper:
        if "PRIVATE" not in comp_upper and "PVT" not in comp_upper and "(P)" not in comp_upper:
            return True, "Public limited company"

    # Multi-garden check (>2 gardens under same company)
    if comp_upper and company_counts.get(comp_upper, 0) > 2:
        return True, f"Multi-garden group ({company_counts[comp_upper]} gardens)"

    return False, ""

def compute_tech_score(row, is_blf, hectares):
    score = 0
    signals = []

    email = clean_text(row.get("email", "")).lower()
    phone = clean_text(row.get("phone", ""))
    hq_city = clean_text(row.get("hq_city", "")).title()

    # Email analysis
    if not email:
        score += 2
        signals.append("No email listed (paper-dominant)")
    elif "@" in email:
        domain = email.split("@")[-1]
        if domain in CONSUMER_EMAIL_DOMAINS:
            score += 4
            signals.append(f"Consumer/ISP email (@{domain})")
        else:
            score += 1
            signals.append(f"Custom domain (@{domain})")

    # Factory / Business model
    if is_blf:
        score += 5
        signals.append("Bought Leaf Factory (daily small-grower leaf friction)")

    # Scale fit for estates
    if not is_blf and hectares is not None:
        if 50 <= hectares <= 350:
            score += 3
            signals.append(f"Mid-tier acreage sweet spot ({hectares} Ha)")
        elif hectares < 50:
            score += 1
            signals.append(f"Small holding ({hectares} Ha)")
        elif hectares > 500:
            score -= 2
            signals.append(f"Large enterprise acreage ({hectares} Ha)")

    # Local proximity
    if hq_city == "Siliguri":
        score += 3
        signals.append("Siliguri HQ (local operating hub)")
    elif hq_city in ["Kolkata", "Jalpaiguri", "Guwahati", "Dibrugarh"]:
        score += 1
        signals.append(f"{hq_city} HQ")

    return score, signals

def load_golden_phones():
    phones = {}
    golden_path = os.path.join(REPO_ROOT, "marketing/outreach/sep-2026-pilot/locked/Contacts Verified Golden List.xlsx")
    if not os.path.exists(golden_path):
        return phones
    try:
        import openpyxl
        wb = openpyxl.load_workbook(golden_path, data_only=True)
        sheet = wb["Sheet1"]
        for i in range(3, sheet.max_row + 1):
            name = sheet.cell(row=i, column=2).value
            phone = sheet.cell(row=i, column=6).value
            if name and phone:
                norm_key = re.sub(r"[^a-z0-9]", "", str(name).lower())
                phones[norm_key] = str(phone).strip()
    except Exception as e:
        print(f"Notice: Could not load golden list phones: {e}")
    return phones

def clean_entity_name(name):
    name = str(name).strip()
    # Remove leading "Ltd." or "M/S" or "M/s."
    name = re.sub(r"^(?:Ltd\.?|M/s\.?|M/S)\s*", "", name, flags=re.IGNORECASE).strip()
    
    # Check if exact halves are duplicated: e.g. "ABC ABC"
    parts = name.split()
    half = len(parts) // 2
    if len(parts) >= 2 and parts[:half] == parts[half:half*2]:
        name = " ".join(parts[:half])
    
    # Check if comma separated halves are duplicated: "ABC, ABC"
    if "," in name:
        subparts = [p.strip() for p in name.split(",") if p.strip()]
        if len(subparts) == 2 and subparts[0].lower() == subparts[1].lower():
            name = subparts[0]
            
    return name.title()

def get_conversational_name(name):
    clean = clean_entity_name(name)
    # Strip formal suffixes for natural sentence flow
    clean = re.sub(r"\s*(?:Pvt\.?\s*Ltd\.?|Private\s*Limited|Co\.?\s*Ltd\.?|Company\s*Limited|Limited|Ltd\.?)\s*$", "", clean, flags=re.IGNORECASE).strip()
    return clean

def generate_engagement_hook(row, is_blf):
    raw_name = clean_text(row.get("name", ""))
    entity_clean = get_conversational_name(raw_name)
    dm = clean_text(row.get("decision_maker", ""))
    district = clean_text(row.get("district", ""))

    # Clean decision maker salutation
    dm_clean = dm
    dm_clean = re.sub(r"^(MR\.|MR|SHRI|DR\.|MRS\.|MS\.)\s*", "", dm_clean, flags=re.IGNORECASE).strip()
    if not dm_clean or dm_clean.upper() in ["-", "N/A", "UNKNOWN"]:
        salutation = "Namaskar,"
    else:
        # If first name is identifiable, use respectful [First Name]-ji
        first_name = dm_clean.split()[0].title()
        salutation = f"Namaskar {first_name}-ji,"

    if is_blf:
        return (
            f"{salutation}\n\n"
            f"During peak leaf intake at {entity_clean}, how does your weighbridge team handle small grower disputes "
            f"over tare weight and moisture or coarse leaf deductions?\n\n"
            f"In {district}, several BLF owners we know were losing 15 to 20 minutes per vehicle because "
            f"deductions were still handwritten or keyed manually into a local computer slip.\n\n"
            f"Does your weighing scale push verified net leaf weights directly to the grower slip and your ledger "
            f"without manual re-entry?"
        )
    else:
        return (
            f"{salutation}\n\n"
            f"When green leaf trucks move from field divisions to the factory at {entity_clean}, what percentage difference "
            f"are you typically seeing between the field scale and factory weighbridge?\n\n"
            f"Most independent estates in {district} tell us they only catch leaf discrepancies and proxy hazira "
            f"10 to 14 days later when the fortnightly muster roll reaches the office.\n\n"
            f"Are you able to see your daily leaf reconciliation and worker count on your phone before the evening shift, "
            f"or is your team still compiling it on paper?"
        )

def process_registry(input_path, state_name, golden_phones=None):
    if golden_phones is None:
        golden_phones = {}
    with open(input_path, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    # Pre-count companies to spot groups
    comp_counts = Counter(
        clean_text(r.get("company", "")).upper()
        for r in rows
        if clean_text(r.get("company", ""))
    )

    filtered_candidates = []
    send_ready_records = []

    for r in rows:
        name = clean_text(r.get("name", ""))
        comp = clean_text(r.get("company", ""))
        entry_type = clean_text(r.get("type", ""))
        is_blf = "Bought Leaf" in entry_type or "BLF" in entry_type
        scale_str = clean_text(r.get("scale", ""))
        hectares = parse_hectares(scale_str)
        phone = clean_text(r.get("phone", ""))
        if not phone and golden_phones:
            norm_name = re.sub(r"[^a-z0-9]", "", name.lower())
            for gk, gv in golden_phones.items():
                if gk in norm_name or norm_name in gk:
                    phone = gv
                    break
        email = clean_text(r.get("email", ""))
        hq_address = clean_text(r.get("hq_address", ""))
        hq_city = clean_text(r.get("hq_city", ""))
        dm = clean_text(r.get("decision_maker", ""))
        district = clean_text(r.get("district", ""))

        corp_flag, corp_reason = is_corporate(comp, name, comp_counts)

        score, signals = compute_tech_score(r, is_blf, hectares)

        # Qualification status
        if corp_flag:
            status = "Disqualified"
            disqual_reason = corp_reason
            tier = "Disqualified"
        elif not is_blf and hectares is not None and hectares > 600:
            status = "Disqualified"
            disqual_reason = f"Acreage too large for independent owner tier ({hectares} Ha)"
            tier = "Disqualified"
        elif score >= 7:
            status = "Qualified"
            disqual_reason = ""
            tier = "Tier 1 (High Leakage / Paper Priority)"
        elif score >= 4:
            status = "Qualified"
            disqual_reason = ""
            tier = "Tier 2 (Qualified Independent)"
        else:
            status = "Uncertain / Backlog"
            disqual_reason = "Low signal density"
            tier = "Tier 3"

        # Channel Allocation
        has_mobile = is_mobile_number(phone)
        has_hq = bool(hq_address and len(hq_address) > 10)
        has_email = bool(email and "@" in email)

        channels = []
        if has_mobile:
            channels.append("WhatsApp / Call")
        if has_hq:
            channels.append("1-Page Postal Dispatch")
        if has_email:
            channels.append("Email")

        if not channels:
            primary_channel = "Requires Phone/Address Recon"
        elif has_mobile:
            primary_channel = "WhatsApp / Phone"
        elif has_hq:
            primary_channel = "1-Page Postal Letter"
        else:
            primary_channel = "Email"

        engagement_hook = generate_engagement_hook(r, is_blf)

        record_summary = {
            "state": state_name,
            "name": name,
            "type": entry_type,
            "is_blf": "Yes" if is_blf else "No",
            "company": comp,
            "district": district,
            "scale_hectares": str(hectares) if hectares else scale_str,
            "decision_maker": dm,
            "phone": clean_phone_number(phone),
            "phone_raw": phone,
            "is_mobile": "Yes" if has_mobile else "No",
            "email": email,
            "hq_city": hq_city,
            "hq_address": hq_address,
            "status": status,
            "tier": tier,
            "tech_score": score,
            "score_signals": " | ".join(signals),
            "disqualification_reason": disqual_reason,
            "primary_channel": primary_channel,
            "available_channels": " + ".join(channels) if channels else "None",
            "engagement_hook": engagement_hook
        }

        filtered_candidates.append(record_summary)

        if status == "Qualified":
            send_ready_records.append(record_summary)

    return filtered_candidates, send_ready_records

def write_csv(filepath, rows, fieldnames):
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    with open(filepath, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

def main():
    print("=== GardenSuite Automated Outreach Filtering & Generation Engine ===")
    os.makedirs(FILTERED_DIR, exist_ok=True)
    os.makedirs(SEND_READY_DIR, exist_ok=True)

    golden_phones = load_golden_phones()
    if golden_phones:
        print(f"Loaded {len(golden_phones)} verified phones from locked Golden List.")

    # Process West Bengal
    print(f"\nProcessing West Bengal Master ({WB_MASTER})...")
    wb_filtered, wb_send_ready = process_registry(WB_MASTER, "West Bengal", golden_phones)
    print(f"  Total WB records: {len(wb_filtered)}")
    print(f"  Qualified send-ready WB records: {len(wb_send_ready)}")

    # Process Assam
    print(f"\nProcessing Assam Master ({ASSAM_MASTER})...")
    assam_filtered, assam_send_ready = process_registry(ASSAM_MASTER, "Assam", golden_phones)
    print(f"  Total Assam records: {len(assam_filtered)}")
    print(f"  Qualified send-ready Assam records: {len(assam_send_ready)}")

    headers = [
        "state", "name", "type", "is_blf", "company", "district", "scale_hectares",
        "decision_maker", "phone", "phone_raw", "is_mobile", "email", "hq_city", "hq_address",
        "status", "tier", "tech_score", "score_signals", "disqualification_reason",
        "primary_channel", "available_channels", "engagement_hook"
    ]
    filtered_headers = headers

    # Write Intermediate Filtered Lists
    wb_filtered_file = os.path.join(FILTERED_DIR, "wb_filtered_candidates.csv")
    write_csv(wb_filtered_file, wb_filtered, filtered_headers)
    print(f"\n[Saved] WB Filtered candidates: {wb_filtered_file}")

    assam_filtered_file = os.path.join(FILTERED_DIR, "assam_filtered_candidates.csv")
    write_csv(assam_filtered_file, assam_filtered, filtered_headers)
    print(f"[Saved] Assam Filtered candidates: {assam_filtered_file}")

    # Write Final Send-Ready Lists
    wb_final_file = os.path.join(SEND_READY_DIR, "wb_final_send_ready.csv")
    write_csv(wb_final_file, wb_send_ready, headers)
    print(f"\n[Saved] WB Final Send-Ready: {wb_final_file} ({len(wb_send_ready)} accounts)")

    # Specialized Sub-Lists for WB
    wb_blf = [r for r in wb_send_ready if r["is_blf"] == "Yes"]
    wb_blf_file = os.path.join(SEND_READY_DIR, "wb_blf_grower_outreach.csv")
    write_csv(wb_blf_file, wb_blf, headers)
    print(f"[Saved] WB Bought Leaf Factories (BLF): {wb_blf_file} ({len(wb_blf)} factories)")

    wb_siliguri = [r for r in wb_send_ready if r["hq_city"].strip().title() == "Siliguri"]
    wb_siliguri_file = os.path.join(SEND_READY_DIR, "wb_siliguri_local_outreach.csv")
    write_csv(wb_siliguri_file, wb_siliguri, headers)
    print(f"[Saved] WB Siliguri HQ Local Outreach: {wb_siliguri_file} ({len(wb_siliguri)} accounts)")

    # Final Assam Send-Ready
    assam_final_file = os.path.join(SEND_READY_DIR, "assam_final_send_ready.csv")
    write_csv(assam_final_file, assam_send_ready, headers)
    print(f"[Saved] Assam Final Send-Ready: {assam_final_file} ({len(assam_send_ready)} accounts)")

    # Summary Statistics
    print("\n=== Pipeline Execution Summary ===")
    print(f"Total Master Records Ingested: {len(wb_filtered) + len(assam_filtered)}")
    print(f"Total Qualified Send-Ready Accounts: {len(wb_send_ready) + len(assam_send_ready)}")
    print(f"  - West Bengal Send-Ready: {len(wb_send_ready)}")
    print(f"    * Bought Leaf Factories (BLF): {len(wb_blf)}")
    print(f"    * Siliguri Local HQ Cluster: {len(wb_siliguri)}")
    print(f"  - Assam Send-Ready: {len(assam_send_ready)}")

if __name__ == "__main__":
    main()
