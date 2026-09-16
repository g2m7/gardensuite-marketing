import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load master rows
with open("marketing/outreach/north-bengal-intelligence/data/North_Bengal_Tea_Master_Registry.json") as f:
    master_rows = json.load(f)

wb = openpyxl.Workbook()

# Setup sheets
ws_data = wb.active
ws_data.title = "Master Registry (396 Units)"
ws_legend = wb.create_sheet(title="Legend & Outreach Guide")

# -------------------------------------------------------------
# PALETTE DEFINITIONS (Clean, executive, subtle)
# -------------------------------------------------------------
FONT_FAMILY = "Segoe UI"

font_header = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

font_title = Font(name=FONT_FAMILY, size=16, bold=True, color="0F172A")
font_subtitle = Font(name=FONT_FAMILY, size=10, color="64748B")

font_body = Font(name=FONT_FAMILY, size=9, color="1E293B")
font_bold = Font(name=FONT_FAMILY, size=9, bold=True, color="0F172A")
font_muted = Font(name=FONT_FAMILY, size=8.5, color="64748B")

# Unit Type Fills & Fonts
fill_estate = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # soft green
font_estate = Font(name=FONT_FAMILY, size=9, bold=True, color="166534")

fill_blf = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid") # soft blue
font_blf = Font(name=FONT_FAMILY, size=9, bold=True, color="1E40AF")

# Region Fills & Fonts
region_styles = {
    "Dooars": (PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="065F46")),
    "Terai": (PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="0369A1")),
    "Darjeeling Hills": (PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="6B21A8")),
    "Uttar Dinajpur": (PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="92400E")),
    "Cooch Behar": (PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="9A3412"))
}

# HQ City Fills & Fonts
hq_styles = {
    "Siliguri": (PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"), Font(name=FONT_FAMILY, size=9, bold=True, color="065F46")),
    "Jalpaiguri": (PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid"), Font(name=FONT_FAMILY, size=9, bold=True, color="0F766E")),
    "Kolkata": (PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="334155")),
    "Islampur": (PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="854D0E")),
    "Darjeeling": (PatternFill(start_color="F5F3FF", end_color="F5F3FF", fill_type="solid"), Font(name=FONT_FAMILY, size=9, color="5B21B6"))
}

thin_border_side = Side(border_style="thin", color="E2E8F0")
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

# -------------------------------------------------------------
# SHEET 1: MASTER REGISTRY
# -------------------------------------------------------------
headers = [
    ("Sl.", 6, "center"),
    ("Unit Name", 32, "left"),
    ("Unit Type", 18, "center"),
    ("Region", 16, "center"),
    ("District", 16, "left"),
    ("Operating Company / Group", 30, "left"),
    ("Key Decision Maker", 22, "left"),
    ("HQ City", 14, "center"),
    ("HQ Address", 36, "left"),
    ("Phone / Mobile", 18, "left"),
    ("Email Address", 28, "left"),
    ("Scale / Capacity", 16, "right"),
    ("Registration / Marks", 22, "left")
]

# Write Header Row
ws_data.row_dimensions[1].height = 28
for col_idx, (h_name, width, align) in enumerate(headers, start=1):
    cell = ws_data.cell(row=1, column=col_idx, value=h_name)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = cell_border
    ws_data.column_dimensions[get_column_letter(col_idx)].width = width

# Write Data Rows
for row_idx, r in enumerate(master_rows, start=2):
    ws_data.row_dimensions[row_idx].height = 22
    is_even = (row_idx % 2 == 0)
    base_row_fill = zebra_fill if is_even else PatternFill(fill_type=None)

    # 1. Sl
    c1 = ws_data.cell(row=row_idx, column=1, value=r["sl"])
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c1.font = font_muted
    c1.fill = base_row_fill
    c1.border = cell_border

    # 2. Name
    c2 = ws_data.cell(row=row_idx, column=2, value=r["name"])
    c2.alignment = Alignment(horizontal="left", vertical="center")
    c2.font = font_bold
    c2.fill = base_row_fill
    c2.border = cell_border

    # 3. Type
    c3 = ws_data.cell(row=row_idx, column=3, value=r["type"])
    c3.alignment = Alignment(horizontal="center", vertical="center")
    if r["type"] == "Organized Estate":
        c3.fill = fill_estate
        c3.font = font_estate
    else:
        c3.fill = fill_blf
        c3.font = font_blf
    c3.border = cell_border

    # 4. Region
    c4 = ws_data.cell(row=row_idx, column=4, value=r["region"])
    c4.alignment = Alignment(horizontal="center", vertical="center")
    r_fill, r_font = region_styles.get(r["region"], (base_row_fill, font_body))
    c4.fill = r_fill
    c4.font = r_font
    c4.border = cell_border

    # 5. District
    c5 = ws_data.cell(row=row_idx, column=5, value=r["district"])
    c5.alignment = Alignment(horizontal="left", vertical="center")
    c5.font = font_body
    c5.fill = base_row_fill
    c5.border = cell_border

    # 6. Company
    c6 = ws_data.cell(row=row_idx, column=6, value=r["company"])
    c6.alignment = Alignment(horizontal="left", vertical="center")
    c6.font = font_body
    c6.fill = base_row_fill
    c6.border = cell_border

    # 7. Decision Maker
    c7 = ws_data.cell(row=row_idx, column=7, value=r["decision_maker"])
    c7.alignment = Alignment(horizontal="left", vertical="center")
    c7.font = font_muted if r["decision_maker"] == "-" else font_body
    c7.fill = base_row_fill
    c7.border = cell_border

    # 8. HQ City
    c8 = ws_data.cell(row=row_idx, column=8, value=r["hq_city"])
    c8.alignment = Alignment(horizontal="center", vertical="center")
    h_fill, h_font = hq_styles.get(r["hq_city"], (base_row_fill, font_body))
    c8.fill = h_fill
    c8.font = h_font
    c8.border = cell_border

    # 9. HQ Address
    c9 = ws_data.cell(row=row_idx, column=9, value=r["hq_address"])
    c9.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    c9.font = font_muted
    c9.fill = base_row_fill
    c9.border = cell_border

    # 10. Phone
    c10 = ws_data.cell(row=row_idx, column=10, value=r["phone"])
    c10.alignment = Alignment(horizontal="left", vertical="center")
    c10.font = font_body if r["phone"] != "-" else font_muted
    c10.fill = base_row_fill
    c10.border = cell_border

    # 11. Email
    c11 = ws_data.cell(row=row_idx, column=11, value=r["email"])
    c11.alignment = Alignment(horizontal="left", vertical="center")
    c11.font = font_body if r["email"] != "-" else font_muted
    c11.fill = base_row_fill
    c11.border = cell_border

    # 12. Scale
    c12 = ws_data.cell(row=row_idx, column=12, value=r["scale"])
    c12.alignment = Alignment(horizontal="right", vertical="center")
    c12.font = font_body
    c12.fill = base_row_fill
    c12.border = cell_border

    # 13. Registration / Marks
    c13 = ws_data.cell(row=row_idx, column=13, value=r["reg_no"])
    c13.alignment = Alignment(horizontal="left", vertical="center")
    c13.font = font_muted
    c13.fill = base_row_fill
    c13.border = cell_border

# Freeze Panes and Auto Filter
ws_data.freeze_panes = "C2" # Freeze header and Sl / Name
ws_data.auto_filter.ref = ws_data.dimensions

# -------------------------------------------------------------
# SHEET 2: LEGEND & OUTREACH GUIDE
# -------------------------------------------------------------
ws_legend.column_dimensions["A"].width = 4
ws_legend.column_dimensions["B"].width = 24
ws_legend.column_dimensions["C"].width = 28
ws_legend.column_dimensions["D"].width = 50

# Title
ws_legend["B2"].value = "North Bengal Tea Master Registry - Legend & Outreach Playbook"
ws_legend["B2"].font = font_title
ws_legend["B3"].value = "Unified operating census of 396 tea estates and standalone bought leaf factories across North Bengal."
ws_legend["B3"].font = font_subtitle

# Summary Metrics Table
ws_legend["B5"].value = "Metric"
ws_legend["C5"].value = "Count"
ws_legend["D5"].value = "Scope & Business Relevance"
for col in ["B", "C", "D"]:
    cell = ws_legend[f"{col}5"]
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = cell_border

metrics_data = [
    ("Total Operating Units", "396", "100% census of commercial tea gardens & processing units in West Bengal"),
    ("Organized Tea Estates", "324", "Captive plantations holding statutory Tea Board registrations & workforce"),
    ("Bought Leaf Factories (BLF)", "72", "Commercial factories processing leaf purchased from small tea growers (STGs)"),
    ("Siliguri Corporate HQs", "57+", "Local commercial headquarters reachable within 30 minutes from Bagdogra"),
    ("Jalpaiguri Local Offices", "16+", "Regional operational hubs situated in the Dooars entrance corridor"),
    ("Kolkata Corporate HQs", "146+", "Boardrooms and executive managing directors of major tea plantation groups"),
    ("Total Tracked Planted Area", "110,000+ Ha", "Combined estate acreage under organized commercial tea cultivation")
]

for idx, (m, c, d) in enumerate(metrics_data, start=6):
    ws_legend[f"B{idx}"].value = m
    ws_legend[f"B{idx}"].font = font_bold
    ws_legend[f"B{idx}"].border = cell_border
    ws_legend[f"C{idx}"].value = c
    ws_legend[f"C{idx}"].font = font_bold
    ws_legend[f"C{idx}"].border = cell_border
    ws_legend[f"D{idx}"].value = d
    ws_legend[f"D{idx}"].font = font_body
    ws_legend[f"D{idx}"].border = cell_border

# Color Coding Legend: Unit Type
ws_legend["B15"].value = "Unit Type Color"
ws_legend["C15"].value = "Classification"
ws_legend["D15"].value = "GardenSuite Product Fit & Outreach Angle"
for col in ["B", "C", "D"]:
    cell = ws_legend[f"{col}15"]
    cell.font = font_header
    cell.fill = fill_header
    cell.border = cell_border

type_legend = [
    ("Soft Green (DCFCE7)", "Organized Estate (324)", "Full captive garden + factory. High-value pitch: Face Recognition Attendance (stops proxy hazira) + Smart Plucking Scale (weighs leaf straight to face) + Factory Production + Daily Cloud MIS."),
    ("Soft Blue (DBEAFE)", "Bought Leaf Factory (BLF) (72)", "Factory without plantation. Buys green leaf from small tea growers. Target pitch: Leaf Collection Weighment Scales + Factory CTC/Orthodox Production ERP + Daily MIS.")
]

for idx, (col_code, name, pitch) in enumerate(type_legend, start=16):
    c_b = ws_legend[f"B{idx}"]
    c_b.value = col_code
    c_b.font = font_estate if "Green" in col_code else font_blf
    c_b.fill = fill_estate if "Green" in col_code else fill_blf
    c_b.border = cell_border

    c_c = ws_legend[f"C{idx}"]
    c_c.value = name
    c_c.font = font_bold
    c_c.border = cell_border

    c_d = ws_legend[f"D{idx}"]
    c_d.value = pitch
    c_d.font = font_body
    c_d.border = cell_border

# Color Coding Legend: HQ Proximity
ws_legend["B20"].value = "HQ City Color"
ws_legend["C20"].value = "Outreach Speed"
ws_legend["D20"].value = "Operational Outreach Protocol"
for col in ["B", "C", "D"]:
    cell = ws_legend[f"{col}20"]
    cell.font = font_header
    cell.fill = fill_header
    cell.border = cell_border

hq_legend = [
    ("Mint Green (D1FAE5)", "Siliguri (Local - Same Day)", "Immediate direct demo route. Decision-makers sit in Sevoke Rd, Hill Cart Rd, or Hakim Para. Dispatch local sales engineer for live tablet demonstration."),
    ("Teal (CCFBF1)", "Jalpaiguri (Regional - 1 Hour)", "Close regional access. Target owners and senior partners in Kadamtalla and Mal Bazar. Schedule morning road trips."),
    ("Slate Blue (F1F5F9)", "Kolkata (Corporate Executive)", "Group headquarters and managing directors. Execute targeted executive email sequence + director phone calls; schedule clustered Kolkata boardroom visits.")
]

for idx, (col_code, name, pitch) in enumerate(hq_legend, start=21):
    c_b = ws_legend[f"B{idx}"]
    c_b.value = col_code
    c_b.font = hq_styles.get(name.split()[0], (None, font_body))[1]
    c_b.fill = hq_styles.get(name.split()[0], (zebra_fill, None))[0]
    c_b.border = cell_border

    c_c = ws_legend[f"C{idx}"]
    c_c.value = name
    c_c.font = font_bold
    c_c.border = cell_border

    c_d = ws_legend[f"D{idx}"]
    c_d.value = pitch
    c_d.font = font_body
    c_d.border = cell_border

# Save Excel
xlsx_path = "marketing/outreach/north-bengal-intelligence/data/North_Bengal_Tea_Master_Registry.xlsx"
wb.save(xlsx_path)
print(f"Saved styled Excel workbook successfully to: {xlsx_path}")
